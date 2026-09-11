"""Build the Test Assistent MCP Server deck reproducing the REAL Cerios
template visual signature:
  - dark navy -> teal diagonal gradient background
  - semi-transparent teal content card with dark text
  - white bold title + bright teal section label
  - thin dark footer bar
  - cerios(R) wordmark bottom-left
"""

from pptx import Presentation
from pptx.util import Pt, Inches, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.oxml.ns import qn
from copy import deepcopy

TEMPLATE = "Cerios AI Community meeting 2025 Q2.pptx"
OUT = "Test-Assistent-MCP-Server.pptx"

NAVY = RGBColor(0x1B, 0x2A, 0x4A)
TEAL = RGBColor(0x6A, 0x9A, 0x8E)
TEAL_BRIGHT = RGBColor(0x5A, 0xBF, 0xA0)
CARD = RGBColor(0x5F, 0x9A, 0x8C)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARKTEXT = RGBColor(0x16, 0x24, 0x3A)

prs = Presentation(TEMPLATE)
for sldId in list(prs.slides._sldIdLst):
    rId = sldId.get(qn("r:id"))
    prs.slides._sldIdLst.remove(sldId)
    if rId is not None:
        prs.part.drop_rel(rId)
SW = prs.slide_width
SH = prs.slide_height


def add_gradient_bg(slide, c1=NAVY, c2=TEAL, angle=2250000):
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.shadow.inherit = False
    rect.line.fill.background()
    fill = rect.fill
    fill.gradient()
    gs = fill.gradient_stops
    gs[0].position = 0.0
    gs[0].color.rgb = c1
    gs[1].position = 1.0
    gs[1].color.rgb = c2
    gf = fill._xPr.find(qn("a:gradFill"))
    lin = gf.find(qn("a:lin"))
    if lin is None:
        from lxml import etree
        lin = etree.SubElement(gf, qn("a:lin"))
    lin.set("ang", str(angle))
    lin.set("scaled", "1")
    rect.shadow.inherit = False
    # send to back
    spTree = slide.shapes._spTree
    spTree.remove(rect._element)
    spTree.insert(2, rect._element)
    return rect


def add_footer(slide):
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SH - Inches(0.32), SW, Inches(0.32))
    bar.fill.solid()
    bar.fill.fore_color.rgb = NAVY
    bar.line.fill.background()
    bar.shadow.inherit = False
    # wordmark
    tb = slide.shapes.add_textbox(Inches(0.5), SH - Inches(0.34), Inches(3.0), Inches(0.3))
    tf = tb.text_frame
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "cerios®"
    r.font.size = Pt(11)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = WHITE
    return bar


def set_alpha(shape, val):
    sf = shape.fill._xPr.find(qn("a:solidFill"))
    srgb = sf.find(qn("a:srgbClr"))
    a = srgb.makeelement(qn("a:alpha"), {"val": str(val)})
    srgb.append(a)


def content_card(slide, top=Inches(1.75), height=Inches(5.0), alpha=82000):
    card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                  Inches(0.5), top, Inches(12.3), height)
    card.fill.solid()
    card.fill.fore_color.rgb = CARD
    card.line.fill.background()
    card.shadow.inherit = False
    set_alpha(card, alpha)
    return card


def title_block(slide, title, label=None):
    if label:
        lb = slide.shapes.add_textbox(Inches(0.6), Inches(0.55), Inches(12), Inches(0.4))
        p = lb.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = label
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.name = "Calibri"
        r.font.color.rgb = TEAL_BRIGHT
    tb = slide.shapes.add_textbox(Inches(0.6), Inches(0.95) if label else Inches(0.6),
                                  Inches(12), Inches(0.9))
    p = tb.text_frame.paragraphs[0]
    r = p.add_run()
    r.text = title
    r.font.size = Pt(32)
    r.font.bold = True
    r.font.name = "Calibri"
    r.font.color.rgb = WHITE
    return tb


def text_in_card(slide, card, paras, x=0.9, y=0.15, w=11.5, h_k=0.7):
    tb = slide.shapes.add_textbox(
        card.left + Inches(x), card.top + Inches(y),
        card.width - Inches(2 * x), card.height - Inches(y + 0.2))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.TOP
    for i, item in enumerate(paras):
        text, opts = item
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.level = opts.get("level", 0)
        run = p.add_run()
        run.text = text
        f = run.font
        f.size = Pt(opts.get("size", 18))
        f.bold = opts.get("bold", False)
        f.italic = opts.get("italic", False)
        f.name = opts.get("font", "Calibri")
        f.color.rgb = opts.get("color", DARKTEXT)
        if "space_after" in opts:
            p.space_after = Pt(opts["space_after"])
    return tb


def code_in_card(slide, card, code_lines, intro=None, code_size=13):
    inner = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   card.left + Inches(0.3), card.top + Inches(0.25),
                                   card.width - Inches(0.6), card.height - Inches(0.5))
    inner.fill.solid()
    inner.fill.fore_color.rgb = RGBColor(0xF4, 0xF7, 0xFA)
    inner.line.color.rgb = RGBColor(0xCC, 0xD6, 0xE2)
    inner.line.width = Pt(0.75)
    inner.shadow.inherit = False
    tb = slide.shapes.add_textbox(
        inner.left + Inches(0.18), inner.top + Inches(0.12),
        inner.width - Inches(0.36), inner.height - Inches(0.24))
    tf = tb.text_frame
    tf.word_wrap = True
    first = True
    if intro:
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = intro
        r.font.size = Pt(14)
        r.font.italic = True
        r.font.name = "Calibri"
        r.font.color.rgb = DARKTEXT
        p.space_after = Pt(8)
        first = False
    for line in code_lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        r = p.add_run()
        r.text = line if line else " "
        r.font.name = "Consolas"
        r.font.size = Pt(code_size)
        r.font.color.rgb = DARKTEXT
        p.space_after = Pt(2)
    return tb


def new_slide():
    return prs.slides.add_slide(prs.slide_masters[0].slide_layouts[0])


# ---------------------------------------------------------------- 1. Titel
s = new_slide()
add_gradient_bg(s)
title_block(s, "Test Assistent MCP Server")
sub = s.shapes.add_textbox(Inches(0.6), Inches(3.3), Inches(12), Inches(2.4))
tf = sub.text_frame
tf.word_wrap = True
for i, (txt, o) in enumerate([
    ("Testtechnieken, heuristieken en testdatageneratie als tools voor je AI-assistent",
     {"size": 22, "italic": True}),
    ("", {"size": 10}),
    ("6 tools  •  stdio  •  Python + FastMCP", {"size": 16, "bold": True}),
    ("Gepubliceerd op Smithery & Glama", {"size": 16}),
]):
    p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
    r = p.add_run()
    r.text = txt
    r.font.size = Pt(o["size"])
    r.font.italic = o.get("italic", False)
    r.font.bold = o.get("bold", False)
    r.font.name = "Calibri"
    r.font.color.rgb = WHITE
    p.space_after = Pt(6)
add_footer(s)

# ---------------------------------------------------------------- 2. Wat is het?
s = new_slide()
add_gradient_bg(s)
title_block(s, "Wat is het?", "Test Assistent MCP Server")
card = content_card(s)
text_in_card(s, card, [
    ("Een MCP-server is een kleine achtergronddienst die de AI-assistent nieuwe "
     "capaciteiten geeft via het Model Context Protocol (stdio + JSON-RPC).",
     {"size": 19, "space_after": 12}),
    ("De Test Assistent MCP Server zet jaren aan testkennis om in kant-en-klare tools:",
     {"size": 19, "bold": True, "space_after": 10}),
    ("Classieke testtechnieken (BVA, equivalentieklassen, decision table, pairwise…)",
     {"size": 18, "space_after": 6}),
    ("Test-heuristieken (SFDPOT, FEW HICCUPPS, RCRCRC, testtours…)",
     {"size": 18, "space_after": 6}),
    ("Property-based testdatageneratie (grenswaarden, willekeurig, met seed)",
     {"size": 18, "space_after": 6}),
    ("Advies en checklists afgeleid uit jouw testcontext",
     {"size": 18, "space_after": 6}),
    ("De server genereert en adviseert alleen testcases — de agent voert ze uit tegen "
     "het systeem onder test (SUT).", {"size": 16, "italic": True}),
])
add_footer(s)

# ---------------------------------------------------------------- 3. Wat kan het?
s = new_slide()
add_gradient_bg(s)
title_block(s, "Wat kan het? — de 6 tools", "Test Assistent MCP Server")
card = content_card(s)
tools = [
    ("catalog_techniques", "Lijst alle classieke testtechnieken"),
    ("catalog_heuristics", "Lijst alle test-heuristieken"),
    ("generate_test_cases", "Testcases voor 7 technieken (BVA, EP, DT, pairwise, ...)"),
    ("generate_test_data", "Ruwe testdata: random rijen of property-based"),
    ("advise_technique", "Aanbeveling van techniek op basis van context"),
    ("checklist_for", "Testchecklist voor een context (bijv. RCRCRC)"),
]
paras = [("Beschikbaar voor de assistent:", {"size": 19, "bold": True, "space_after": 8})]
for name, desc in tools:
    paras.append((f"{name}  —  {desc}", {"size": 17, "space_after": 6,
                                          "font": "Consolas"}))
text_in_card(s, card, paras)
add_footer(s)

# ---------------------------------------------------------------- 4. Hoe werkt het?
s = new_slide()
add_gradient_bg(s)
title_block(s, "Hoe werkt het?", "Test Assistent MCP Server")
card = content_card(s)
code_in_card(s, card, [
    "OpenCode (AI-assistent)",
    "      |  tool-aanroep (JSON-RPC)",
    "      v",
    "Test Assistent MCP Server",
    "  - 6 tools",
    "  - lokale kennisbank (JSON)",
    "      |",
    "      v",
    "Systeem onder test (SUT)",
    "  <- de agent voert de testgevallen uit",
], intro="Werkwijze: 1) assistent vraagt techniek/testdata op — 2) server levert "
         "gestructureerde testgevallen — 3) agent past ze toe op de SUT.")
add_footer(s)

# ---------------------------------------------------------------- 5. FO voorbeeld
s = new_slide()
add_gradient_bg(s)
title_block(s, "Voorbeeld: Functionele Omschrijving (FO)", "Test Assistent MCP Server")
card = content_card(s)
code_in_card(s, card, [
    "FO-FR014  —  Leeftijdscontrole bij registratie",
    "",
    "Als geregistreerde gebruiker wil ik bij registratie mijn leeftijd opgeven,",
    "zodat alleen bezoekers van 18 jaar of ouder een account kunnen aanmaken.",
    "",
    "•  Leeftijdsveld accepteert een geheel getal tussen 0 en 120.",
    "•  Bij leeftijd < 18 wordt registratie geweigerd met de melding",
    "      'Je moet 18 jaar of ouder zijn.'",
    "•  Bij leeftijd >= 18 wordt het account aangemaakt.",
    "•  Een niet-numerieke of ontbrekende leeftijd geeft een validatiefout.",
], intro="De input: een functionele requirement uit de specificatie.")
add_footer(s)

# ---------------------------------------------------------------- 6. FO -> BDD (ingekort)
s = new_slide()
add_gradient_bg(s)
title_block(s, "Van FO naar BDD-testscenario's", "Test Assistent MCP Server")
card = content_card(s)
code_in_card(s, card, [
    "Functionaliteit: Registratie met leeftijdscontrole",
    "",
    "  Scenario: Minderjarige wordt geweigerd (grens 17)",
    "    Gegeven een bezoeker vult leeftijd 17 in",
    "    Als de registratie wordt verzonden",
    "    Dan verschijnt 'Je moet 18 jaar of ouder zijn.'",
    "    En wordt er geen account aangemaakt",
    "",
    "  Scenario: 18-jarige mag registreren (grens inklusief)",
    "    Gegeven een bezoeker vult leeftijd 18 in",
    "    Als de registratie wordt verzonden",
    "    Dan wordt het account aangemaakt",
    "",
    "  Scenario: Waarde boven bereik wordt geweigerd (121)",
    "    Gegeven een bezoeker vult leeftijd 121 in",
    "    Als de registratie wordt verzonden",
    "    Dan verschijnt een validatiefout over het bereik 0-120",
], intro="De server levert de grenswaarden (BVA) — de agent zet ze om in Gherkin.",
code_size=12)
add_footer(s)

# ---------------------------------------------------------------- 7. Samenvatting
s = new_slide()
add_gradient_bg(s)
title_block(s, "Samenvatting", "Test Assistent MCP Server")
card = content_card(s)
text_in_card(s, card, [
    ("Testkennis altijd binnen handbereik van de assistent — geen expertise nodig.",
     {"size": 19, "space_after": 10}),
    ("Slimme testdatageneratie (grenswaarden, random, seed) bespaart handwerk.",
     {"size": 19, "space_after": 10}),
    ("FO -> BDD scenario's gaat in seconden, met de juiste grenswaarden.",
     {"size": 19, "space_after": 10}),
    ("Herbruikbare checklists en techniek-advies per context.",
     {"size": 19, "space_after": 10}),
    ("Gepubliceerd op Smithery & Glama  •  claimed & build-gate groen",
     {"size": 17, "bold": True}),
])
add_footer(s)

prs.save(OUT)
print("Saved:", OUT, "slides:", len(prs.slides._sldIdLst))
